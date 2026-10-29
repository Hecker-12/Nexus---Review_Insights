"""File loader and column detector for dropbox module."""

import io
import os
import pandas as pd
from typing import Optional, Tuple
from openpyxl import load_workbook

from src.dropbox.errors import (
    EmptyFileError,
    InvalidFileTypeError,
    FileSizeLimitError,
    NoTextColumnError,
    ColumnNotFoundError,
    SheetLimitError,
)

MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_ROWS = 3000
MAX_COLS = 50


def load_file(
    file_bytes: bytes,
    filename: str,
    text_column: Optional[str] = None
) -> Tuple[pd.DataFrame, str, Optional[str], int, bool]:
    """Load and validate an uploaded CSV or XLSX file.

    Returns:
        (df, detected_text_col, detected_rating_col, rows_in, sampled)
    """
    if not file_bytes or len(file_bytes) == 0:
        raise EmptyFileError("Uploaded file is empty (0 bytes).")

    if len(file_bytes) > MAX_FILE_BYTES:
        size_mb = len(file_bytes) / (1024 * 1024)
        raise FileSizeLimitError(
            f"File size ({size_mb:.2f} MB) exceeds maximum allowed limit of 5.00 MB."
        )

    _, ext = os.path.splitext(filename.lower())
    if ext not in [".csv", ".xlsx"]:
        raise InvalidFileTypeError(
            f"Unsupported file format '{ext}'. Only .csv and .xlsx files are supported."
        )

    buffer = io.BytesIO(file_bytes)

    if ext == ".csv":
        df = None
        for enc in ["utf-8-sig", "latin-1"]:
            try:
                buffer.seek(0)
                df = pd.read_csv(buffer, encoding=enc)
                break
            except Exception:
                continue

        if df is None:
            raise EmptyFileError("Failed to decode CSV with utf-8-sig or latin-1 encoding.")

    else:  # .xlsx
        # Pre-validate row/column caps with openpyxl in read_only mode
        try:
            buffer.seek(0)
            wb = load_workbook(buffer, read_only=True, data_only=True)
            sheet = wb.active
            rows_iter = sheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)

            if header_row is None:
                wb.close()
                raise EmptyFileError("Uploaded Excel file contains no header row.")

            num_cols = len([c for c in header_row if c is not None])
            if num_cols > MAX_COLS:
                wb.close()
                raise SheetLimitError(
                    f"Excel file has {num_cols} columns, exceeding the maximum allowed limit of {MAX_COLS} columns."
                )

            data_row_count = 0
            for _ in rows_iter:
                data_row_count += 1
                if data_row_count > MAX_ROWS:
                    wb.close()
                    raise SheetLimitError(
                        f"Excel file exceeds maximum allowed limit of {MAX_ROWS} data rows."
                    )
            wb.close()

            buffer.seek(0)
            df = pd.read_excel(buffer)
        except (SheetLimitError, EmptyFileError):
            raise
        except Exception as e:
            raise EmptyFileError(f"Failed to parse Excel file: {e}")

    rows_in = len(df)
    if rows_in == 0 or len(df.columns) == 0:
        raise EmptyFileError("Uploaded file contains no rows or columns.")

    sampled = False
    if rows_in > MAX_ROWS:
        df = df.sample(n=MAX_ROWS, random_state=42).reset_index(drop=True)
        sampled = True

    # Detect or validate text column
    if text_column is not None and text_column.strip():
        req_col = text_column.strip()
        if req_col not in df.columns:
            raise ColumnNotFoundError(
                f"Specified text column '{req_col}' not found in columns: {list(df.columns)}"
            )
        detected_text_col = req_col
    else:
        best_col = None
        max_mean_len = -1.0
        for col in df.columns:
            # Exclude purely numeric columns from automatic text detection
            if pd.api.types.is_numeric_dtype(df[col]):
                continue
            non_null_share = df[col].notna().mean()
            if non_null_share >= 0.5:
                # Calculate mean length of non-null strings
                mean_len = df[col].dropna().astype(str).str.len().mean()
                if mean_len > max_mean_len:
                    max_mean_len = mean_len
                    best_col = col

        if best_col is None or max_mean_len <= 0:
            raise NoTextColumnError(
                "Could not automatically detect a text column. "
                "Please specify a text column name."
            )
        detected_text_col = best_col

    # Detect optional rating column
    detected_rating_col = None
    candidate_rating_cols = []
    for col in df.columns:
        if col == detected_text_col:
            continue
        col_lower = str(col).lower()
        if col_lower in ["id", "review_id", "index"] or col_lower.startswith("unnamed"):
            continue
        numeric_series = pd.to_numeric(df[col], errors="coerce")
        valid_numeric = numeric_series.dropna()
        if len(valid_numeric) >= 0.5 * len(df) and len(valid_numeric) > 0:
            in_range_share = (valid_numeric.isin([1, 2, 3, 4, 5]) | ((valid_numeric >= 1.0) & (valid_numeric <= 5.0))).mean()
            if in_range_share >= 0.90:
                candidate_rating_cols.append(col)

    if candidate_rating_cols:
        rating_named = [c for c in candidate_rating_cols if any(k in str(c).lower() for k in ["rat", "star", "score"])]
        if rating_named:
            detected_rating_col = rating_named[0]
        else:
            detected_rating_col = candidate_rating_cols[0]

    return df, detected_text_col, detected_rating_col, rows_in, sampled
